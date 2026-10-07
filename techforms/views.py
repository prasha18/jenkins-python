import logging
import json
from datetime import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.http import JsonResponse, HttpResponse
from django.views.decorators.csrf import csrf_exempt, csrf_protect
from django.core.files.storage import default_storage
from django.contrib import messages
from .forms import InnovationVoucherFormStep1, InnovationVoucherFormStep2, InnovationVoucherFormStep3, InnovationVoucherFormStep4
from .forms import PatentApplicationFormStep1, PatentApplicationFormStep2, PatentApplicationFormStep3
from .forms import PriorArtSearchFormStep1, PriorArtSearchFormStep2
from .forms import TrademarkApplicationFormStep1, TrademarkApplicationFormStep2
from .forms import TRLAssessmentForm, ExpressInterestForm
from .models import InnovationVoucherApplication, PatentApplication, PriorArtSearch, TrademarkApplication, TRLAssessment
from .models import State, District
from django.contrib.contenttypes.models import ContentType
from tech_repo.models import Inventor, State, District, Tech_Director
import csv
from docx import Document
from django.utils import timezone
from .forms import PatentForm
from tech_website.models import Sector



logger = logging.getLogger(__name__)

def get_districts(request, state_id):
    try:
        districts = District.objects.filter(state_id=state_id).values('id', 'name')
        return JsonResponse(list(districts), safe=False)
    except Exception as e:
        logger.error(f"Error fetching districts: {str(e)}")
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
    

@csrf_protect
def patent(request):
    states = State.objects.all()
    sectors = Sector.objects.all()
    if request.method == 'POST':
        form = PatentForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return JsonResponse({'status': 'success'})
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    return render(request, 'tech_forms/patent.html', {'states': states, 'sectors': sectors})

@csrf_protect
def express_interest(request):
    if request.method == 'POST':
        form = ExpressInterestForm(request.POST)
        if form.is_valid():
            form.save()
            return JsonResponse({'status': 'success'})
        else:
            return JsonResponse({'status': 'error', 'errors': form.errors}, status=400)
    
    patent_title = request.GET.get('patent_title', '')
    unique_id = request.GET.get('unique_id', '')  
    if unique_id:
        patent_title = f"{patent_title} - {unique_id}"  
    return render(request, 'tech_forms/express-interest.html', {'patent_title': patent_title})


def trl_assessment_form(request):
    try:
        logger.debug("Received request for trl_assessment_form: %s, Data: %s", request.method, request.POST)
        if request.method == 'POST':
            form = TRLAssessmentForm(request.POST)
            if form.is_valid():
                application = form.save()
                logger.info("TRL Assessment Form submitted successfully by %s", application.inventor_name)
                return JsonResponse({'status': 'success', 'message': 'Form Submitted Successfully'})
            else:
                logger.error("TRL Assessment Form validation failed: %s", form.errors)
                return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
        else:
            form = TRLAssessmentForm()
            return render(request, 'tech_forms/trl-assessmentform.html', {'form': form})
    except Exception as e:
        logger.error("Error in trl_assessment_form: %s", str(e))
        return JsonResponse({'status': 'error', 'errors': {'general': [str(e)]}}, status=500)

import logging
import json
from django.shortcuts import render
from django.http import JsonResponse
from django.contrib.contenttypes.models import ContentType
from django.db import transaction

from .forms import (
    InnovationVoucherFormStep1, 
    InnovationVoucherFormStep2, 
    InnovationVoucherFormStep3, 
    InnovationVoucherFormStep4
)
from .models import InnovationVoucherApplication
from tech_repo.models import State, District, Tech_Director

logger = logging.getLogger(__name__)

def innovation_voucherform(request):
    if request.method == 'POST':
        step = int(request.POST.get('step', 0))
        if step == 3:  
            logger.debug("Submitted data for innovation voucher form: %s", request.POST)
            
            try:
                with transaction.atomic():
                    form0 = InnovationVoucherFormStep1(request.POST, request.FILES)
                    form1 = InnovationVoucherFormStep2(request.POST, request.FILES)
                    form2 = InnovationVoucherFormStep3(request.POST, request.FILES)
                    form3 = InnovationVoucherFormStep4(request.POST, request.FILES)

                    if all([form0.is_valid(), form1.is_valid(), form2.is_valid(), form3.is_valid()]):
                        data = {}
                        data.update(form0.cleaned_data)
                        data.update(form1.cleaned_data)
                        data.update(form2.cleaned_data)
                        data.update(form3.cleaned_data)

                        # Handle director details
                        director_count = data['director_count']
                        director_details = []
                        
                        for i in range(1, director_count + 1):
                            dir_dict = {
                                'name': request.POST.get(f'director_name_{i}', ''),
                                'designation': request.POST.get(f'director_designation_{i}', ''),
                                'dob': request.POST.get(f'director_dob_{i}', ''),
                                'mobile': request.POST.get(f'director_mobile_{i}', ''),
                                'email': request.POST.get(f'director_email_{i}', ''),
                                'aadhar': request.FILES.get(f'director_aadhar_{i}')
                            }
                            
                            # Validate director data
                            if not all([dir_dict['name'].strip(), dir_dict['designation'].strip(), 
                                       dir_dict['dob'], dir_dict['mobile'].strip(), 
                                       dir_dict['email'].strip(), dir_dict['aadhar']]):
                                return JsonResponse({
                                    'status': 'error', 
                                    'errors': {'director_details': [f'All fields are required for director {i}']}
                                }, status=400)
                            
                            director_details.append(dir_dict)

                        # Get state and district IDs
                        state_id = data['state'].id if hasattr(data['state'], 'id') else data['state']
                        district_id = data['district'].id if hasattr(data['district'], 'id') else data['district']

                        # Validate cost calculations
                        total_estimated = float(data.get('total_estimated_cost', 0))
                        total_ivp = float(data.get('total_ivp_contribution', 0))
                        total_applicant = float(data.get('total_applicant_contribution', 0))
                        
                        if abs(total_estimated - (total_ivp + total_applicant)) > 0.01:
                            return JsonResponse({
                                'status': 'error', 
                                'errors': {'total_estimated_cost': ['Total estimated cost must match the sum of IVP and applicant contributions.']}
                            }, status=400)

                        # Create the main application
                        app = InnovationVoucherApplication(
                            startup_name=data['startup_name'],
                            dpiit_number=data['dpiit_number'],
                            address=data['address'],
                            country=data['country'],
                            state_id=state_id,  
                            district_id=district_id,  
                            email=data['email'],
                            contact_number=data['contact_number'],
                            company_classification=data['company_classification'],
                            date_of_registration=data['date_of_registration'],
                            uam_certificate=data['uam_certificate'],
                            turnover=data.get('turnover'),
                            website=data.get('website', ''),
                            director_count=director_count,
                            voucher_type=data['voucher_type'],
                            project_title=data['project_title'],
                            product_name=data['product_name'],
                            problem_statement=data['problem_statement'],
                            solution_proposed=data['solution_proposed'],
                            detailed_description=data['detailed_description'],
                            invention_upgrade=data['invention_upgrade'],
                            scope=data['scope'],
                            need=data['need'],
                            competitive_advantage=data['competitive_advantage'],
                            social_impact=data['social_impact'],
                            collaboration=data['collaboration'],
                            benefits=data['benefits'],
                            price_advantage=data['price_advantage'],
                            project_cost=data['project_cost'],
                            target_dates=data['target_dates'],
                            ppt_upload=data['ppt_upload'],
                            youtube_link=data.get('youtube_link', ''),
                            outsourcing_charges=data.get('outsourcing_charges', 0),
                            outsourcing_ivp=data.get('outsourcing_ivp', 0),
                            outsourcing_applicant=data.get('outsourcing_applicant', 0),
                            raw_materials=data.get('raw_materials', 0),
                            raw_ivp=data.get('raw_ivp', 0),
                            raw_applicant=data.get('raw_applicant', 0),
                            fabrication_charges=data.get('fabrication_charges', 0),
                            fabrication_ivp=data.get('fabrication_ivp', 0),
                            fabrication_applicant=data.get('fabrication_applicant', 0),
                            ipr_patent=data.get('ipr_patent', 0),
                            ipr_ivp=data.get('ipr_ivp', 0),
                            ipr_applicant=data.get('ipr_applicant', 0),
                            testing_validation=data.get('testing_validation', 0),
                            testing_ivp=data.get('testing_ivp', 0),
                            testing_applicant=data.get('testing_applicant', 0),
                            commercialization_support=data.get('commercialization_support', 0),
                            commercialization_ivp=data.get('commercialization_ivp', 0),
                            commercialization_applicant=data.get('commercialization_applicant', 0),
                            other1_activity=data.get('other1_activity', ''),
                            other1_cost=data.get('other1_cost', 0),
                            other1_ivp=data.get('other1_ivp', 0),
                            other1_applicant=data.get('other1_applicant', 0),
                            other2_activity=data.get('other2_activity', ''),
                            other2_cost=data.get('other2_cost', 0),
                            other2_ivp=data.get('other2_ivp', 0),
                            other2_applicant=data.get('other2_applicant', 0),
                            total_estimated_cost=data.get('total_estimated_cost', 0),
                            total_ivp_contribution=data.get('total_ivp_contribution', 0),
                            total_applicant_contribution=data.get('total_applicant_contribution', 0),
                            request_knowledge_partner=data.get('request_knowledge_partner', False),
                        )
                        app.save()

                        # Create director records using Tech_Director model
                        content_type = ContentType.objects.get_for_model(InnovationVoucherApplication)
                        
                        for director_data in director_details:
                            Tech_Director.objects.create(
                                content_type=content_type,
                                object_id=app.pk,
                                name=director_data['name'],
                                designation=director_data['designation'],
                                dob=director_data['dob'],
                                aadhaar=director_data['aadhar'],
                                mobile=director_data['mobile'],
                                email=director_data['email']
                            )

                        logger.info("Innovation Voucher Application submitted successfully by %s (ID: %s)", app.startup_name, app.pk)
                        return JsonResponse({'status': 'success', 'message': 'Form Submitted Successfully'})
                        
                    else:
                        errors = {}
                        if not form0.is_valid():
                            errors['step0'] = form0.errors.as_json()
                        if not form1.is_valid():
                            errors['step1'] = form1.errors.as_json()
                        if not form2.is_valid():
                            errors['step2'] = form2.errors.as_json()
                        if not form3.is_valid():
                            errors['step3'] = form3.errors.as_json()
                        logger.error("Innovation Voucher Form validation failed: %s", errors)
                        return JsonResponse({'status': 'error', 'errors': json.dumps(errors)}, status=400)
                        
            except Exception as e:
                logger.error("Error processing innovation voucher form: %s", str(e))
                return JsonResponse({'status': 'error', 'errors': f'Error processing form: {str(e)}'}, status=500)
        else:
            return JsonResponse({'status': 'error', 'message': 'Invalid step'}, status=400)
    else: 
        states = State.objects.all()
        return render(request, 'tech_forms/innovation-voucherform.html', {'states': states})

def debug_session(request):
    """Debug endpoint to check session contents"""
    logger.debug("Debug session contents: %s", dict(request.session.items()))
    return JsonResponse({'session': dict(request.session.items())})

def patent_application_form(request):
    try:
        logger.debug("Received request for patent_application_form: %s", request.method)
        if request.method == 'POST':
            step = int(request.POST.get('step', 0))
            logger.debug("Processing step: %d with data: %s", step, request.POST)
            if step == 0:
                form = PatentApplicationFormStep1(request.POST, request.FILES)
                if form.is_valid():
                    application = form.save()  # Save the initial data
                    inventor_count = form.cleaned_data.get('inventor_count', 0)
                    content_type = ContentType.objects.get_for_model(PatentApplication)
                    
                    # Save inventors
                    for inventor_data in form.cleaned_data.get('inventors', []):
                        Inventor.objects.create(
                            content_type=content_type,
                            object_id=application.pk,
                            name=inventor_data['name'],
                            designation=inventor_data['designation'],
                            dob=inventor_data['dob'],
                            aadhaar=inventor_data['aadhaar'],
                            mobile=inventor_data['mobile'],
                            email=inventor_data['email']
                        )
                    
                    # Store the application ID and step 0 data in the session
                    request.session['patent_application_id'] = application.pk
                    request.session['patent_data'] = {
                        'startup_name': application.startup_name,
                        'dpiit_number': application.dpiit_number,
                        'address': application.address,
                        'country': application.country,
                        'state': application.state.id if application.state else None,
                        'district': application.district.id if application.district else None,
                        'email': application.email,
                        'contact_number': application.contact_number,
                        'company_classification': application.company_classification,
                        'date_of_registration': application.date_of_registration.isoformat() if application.date_of_registration else None,
                        'uam_certificate': application.uam_certificate.name if application.uam_certificate else '',
                        'website': application.website or '',
                        'inventor_count': application.inventor_count,
                        'inventors': [{
                            'name': inventor_data['name'],
                            'designation': inventor_data['designation'],
                            'dob': inventor_data['dob'],
                            'aadhaar': inventor_data['aadhaar'].name if inventor_data['aadhaar'] else '',
                            'mobile': inventor_data['mobile'],
                            'email': inventor_data['email']
                        } for inventor_data in form.cleaned_data.get('inventors', [])]
                    }
                    request.session.modified = True
                    return JsonResponse({'status': 'success', 'next_step': 1})
                else:
                    logger.error("Patent Application Form Step 1 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
            elif step == 1:
                application_id = request.session.get('patent_application_id')
                if not application_id:
                    logger.error("No application ID found in session for step 1")
                    return JsonResponse({'status': 'error', 'message': 'Session data missing. Please start over.'}, status=400)
                
                try:
                    application = PatentApplication.objects.get(pk=application_id)
                except PatentApplication.DoesNotExist:
                    logger.error("PatentApplication with ID %s not found", application_id)
                    return JsonResponse({'status': 'error', 'message': 'Application not found. Please start over.'}, status=400)
                
                form = PatentApplicationFormStep2(request.POST, request.FILES, instance=application)
                if form.is_valid():
                    application = form.save()
                    return JsonResponse({'status': 'success', 'next_step': 2})
                else:
                    logger.error("Patent Application Form Step 2 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
            elif step == 2:
                application_id = request.session.get('patent_application_id')
                if not application_id:
                    logger.error("No application ID found in session for step 2")
                    return JsonResponse({'status': 'error', 'message': 'Session data missing. Please start over.'}, status=400)
                
                try:
                    application = PatentApplication.objects.get(pk=application_id)
                except PatentApplication.DoesNotExist:
                    logger.error("PatentApplication with ID %s not found", application_id)
                    return JsonResponse({'status': 'error', 'message': 'Application not found. Please start over.'}, status=400)
                
                form = PatentApplicationFormStep3(request.POST, request.FILES, instance=application)
                if form.is_valid():
                    application = form.save()
                    logger.info("Patent Application Form submitted successfully by %s", application.startup_name)
                    del request.session['patent_application_id']
                    del request.session['patent_data']
                    return JsonResponse({'status': 'success', 'message': 'Form Submitted Successfully'})
                else:
                    logger.error("Patent Application Form Step 3 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
        else:
            states = State.objects.all()
            return render(request, 'tech_forms/patent-applicationnew.html', {'states': states})
    except Exception as e:
        logger.error("Error in patent_application_form: %s", str(e))
        return JsonResponse({'status': 'error', 'errors': str(e)}, status=500)



def patentapplication_export_data(request):
    if request.method == 'POST':
        export_option = request.POST.get('export_option')
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        format_type = request.POST.get('format', 'csv')
        queryset = PatentApplication.objects.all()

        if export_option == 'range':
            if not (start_date and end_date):
                return render(request, 'admin/techforms/patentapplication/export-data.html', {'error': 'Please select both Start Date and End Date for date range export.'})
            queryset = queryset.filter(date_of_registration__gte=start_date, date_of_registration__lte=end_date)
        elif export_option != 'all':
            return render(request, 'admin/techforms/patentapplication/export-data.html', {'error': 'Please select a valid export option.'})

        if format_type == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="patent_applications.csv"'
            writer = csv.writer(response)
            writer.writerow([
                'Startup Name', 'DPIIT Number', 'Date of Registration', 'Email', 'Contact Number',
                'Company Classification', 'Inventor Count', 'Invention Category', 'Stage of Development',
                'Test Status', 'Patent Status', 'Patent Date', 'Patent Number', 'Created At',
                'Address', 'Country', 'State', 'District', 'Inventors'
            ])
            for app in queryset:
                content_type = ContentType.objects.get_for_model(PatentApplication)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                inventors_str = '\n'.join([f"Inventor {i+1}: Name-{inv.name} - Designation-{inv.designation} - DOB-{inv.dob} - Mobile-{inv.mobile} - Email-{inv.email}" for i, inv in enumerate(inventors)])
                writer.writerow([
                    app.startup_name, app.dpiit_number, app.date_of_registration, app.email,
                    app.contact_number, app.company_classification, app.inventor_count,
                    app.invention_category, app.stage_of_development, app.test_status,
                    app.patent_status, app.patent_date, app.patent_number, app.created_at,
                    app.address, app.country, app.state.name if app.state else '',
                    app.district.name if app.district else '', inventors_str
                ])
            return response
        elif format_type == 'word':
            doc = Document()
            for app in queryset:
                content_type = ContentType.objects.get_for_model(PatentApplication)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                doc.add_paragraph(f"Startup name: {app.startup_name}")
                doc.add_paragraph(f"Dpiit number: {app.dpiit_number}")
                doc.add_paragraph(f"Address: {app.address}")
                doc.add_paragraph(f"Country: {app.country}")
                doc.add_paragraph(f"State: {app.state.name if app.state else 'N/A'}")
                doc.add_paragraph(f"District: {app.district.name if app.district else 'N/A'}")
                doc.add_paragraph(f"Email: {app.email}")
                doc.add_paragraph(f"Contact number: {app.contact_number}")
                doc.add_paragraph(f"Company classification: {app.company_classification}")
                doc.add_paragraph(f"Date of registration: {app.date_of_registration}")
                doc.add_paragraph(f"Invention category: {app.invention_category}")
                doc.add_paragraph(f"Stage of development: {app.stage_of_development}")
                doc.add_paragraph(f"Test status: {app.test_status}")
                doc.add_paragraph(f"Patent status: {app.patent_status}")
                doc.add_paragraph(f"Patent date: {app.patent_date}")
                doc.add_paragraph(f"Patent number: {app.patent_number}")
                doc.add_paragraph(f"Created at: {app.created_at}")
                for i, inv in enumerate(inventors, 1):
                    doc.add_paragraph(f"Inventor {i} details :")
                    doc.add_paragraph(f"Name: {inv.name}")
                    doc.add_paragraph(f"Designation: {inv.designation}")
                    doc.add_paragraph(f"Dob: {inv.dob}")
                    doc.add_paragraph(f"Mobile: {inv.mobile}")
                    doc.add_paragraph(f"Email: {inv.email}")
                doc.add_paragraph('')
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = 'attachment; filename="patent_applications.docx"'
            doc.save(response)
            return response

    return render(request, 'admin/techforms/patentapplication/export-data.html', {})


def innovationapplication_export_data(request):
    if request.method == 'POST':
        export_option = request.POST.get('export_option')
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        format_type = request.POST.get('format', 'csv')
        queryset = InnovationVoucherApplication.objects.all()

        if export_option == 'range':
            if not (start_date and end_date):
                return render(request, 'admin/techforms/innovationapplication/export-data.html', 
                            {'error': 'Please select both Start Date and End Date for date range export.'})
            queryset = queryset.filter(date_of_registration__gte=start_date, date_of_registration__lte=end_date)
        elif export_option != 'all':
            return render(request, 'admin/techforms/innovationvoucherapplication/export-data.html', 
                        {'error': 'Please select a valid export option.'})

        if format_type == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="innovation_voucher_applications.csv"'
            writer = csv.writer(response)
            
            # CSV Headers
            writer.writerow([
                'Startup Name', 'DPIIT Number', 'Date of Registration', 'Email', 'Contact Number',
                'Company Classification', 'Director Count', 'Voucher Type', 'Project Title', 'Product Name',
                'Address', 'Country', 'State', 'District', 'Website', 'Turnover',
                'Problem Statement', 'Solution Proposed', 'Detailed Description', 'Invention Upgrade',
                'Scope', 'Need', 'Competitive Advantage', 'Social Impact', 'Collaboration', 'Benefits',
                'Price Advantage', 'Project Cost', 'Target Dates', 'YouTube Link',
                'Outsourcing Charges', 'Outsourcing IVP', 'Outsourcing Applicant',
                'Raw Materials', 'Raw IVP', 'Raw Applicant',
                'Fabrication Charges', 'Fabrication IVP', 'Fabrication Applicant',
                'IPR Patent', 'IPR IVP', 'IPR Applicant',
                'Testing Validation', 'Testing IVP', 'Testing Applicant',
                'Commercialization Support', 'Commercialization IVP', 'Commercialization Applicant',
                'Other1 Activity', 'Other1 Cost', 'Other1 IVP', 'Other1 Applicant',
                'Other2 Activity', 'Other2 Cost', 'Other2 IVP', 'Other2 Applicant',
                'Total Estimated Cost', 'Total IVP Contribution', 'Total Applicant Contribution',
                'Request Knowledge Partner', 'Created At', 'Directors'
            ])
            
            for app in queryset:
                content_type = ContentType.objects.get_for_model(InnovationVoucherApplication)
                directors = Tech_Director.objects.filter(content_type=content_type, object_id=app.id)
                directors_str = '\n'.join([
                    f"Director {i+1}: Name-{dir.name} - Designation-{dir.designation} - DOB-{dir.dob} - Mobile-{dir.mobile} - Email-{dir.email}" 
                    for i, dir in enumerate(directors)
                ])
                
                writer.writerow([
                    app.startup_name, app.dpiit_number, app.date_of_registration, app.email,
                    app.contact_number, app.company_classification, app.director_count, 
                    app.get_voucher_type_display() if hasattr(app, 'get_voucher_type_display') else app.voucher_type,
                    app.project_title, app.product_name,
                    app.address, app.country, 
                    app.state.name if app.state else '', app.district.name if app.district else '',
                    app.website or '', app.turnover or '',
                    app.problem_statement, app.solution_proposed, app.detailed_description, app.invention_upgrade,
                    app.scope, app.need, app.competitive_advantage, app.social_impact, app.collaboration, app.benefits,
                    app.price_advantage, app.project_cost, app.target_dates, app.youtube_link or '',
                    app.outsourcing_charges or '', app.outsourcing_ivp or '', app.outsourcing_applicant or '',
                    app.raw_materials or '', app.raw_ivp or '', app.raw_applicant or '',
                    app.fabrication_charges or '', app.fabrication_ivp or '', app.fabrication_applicant or '',
                    app.ipr_patent or '', app.ipr_ivp or '', app.ipr_applicant or '',
                    app.testing_validation or '', app.testing_ivp or '', app.testing_applicant or '',
                    app.commercialization_support or '', app.commercialization_ivp or '', app.commercialization_applicant or '',
                    app.other1_activity or '', app.other1_cost or '', app.other1_ivp or '', app.other1_applicant or '',
                    app.other2_activity or '', app.other2_cost or '', app.other2_ivp or '', app.other2_applicant or '',
                    app.total_estimated_cost or '', app.total_ivp_contribution or '', app.total_applicant_contribution or '',
                    'Yes' if app.request_knowledge_partner else 'No', app.created_at, directors_str
                ])
            return response
            
        elif format_type == 'word':
            doc = Document()
            doc.add_heading('Innovation Voucher Applications Export', 0)
            
            for app in queryset:
                content_type = ContentType.objects.get_for_model(InnovationVoucherApplication)
                directors = Tech_Director.objects.filter(content_type=content_type, object_id=app.id)
                
                # Application Header
                doc.add_heading(f'Application: {app.startup_name}', level=1)
                
                # Basic Information
                doc.add_heading('Basic Information', level=2)
                doc.add_paragraph(f"Startup Name: {app.startup_name}")
                doc.add_paragraph(f"DPIIT Number: {app.dpiit_number}")
                doc.add_paragraph(f"Address: {app.address}")
                doc.add_paragraph(f"Country: {app.country}")
                doc.add_paragraph(f"State: {app.state.name if app.state else 'N/A'}")
                doc.add_paragraph(f"District: {app.district.name if app.district else 'N/A'}")
                doc.add_paragraph(f"Email: {app.email}")
                doc.add_paragraph(f"Contact Number: {app.contact_number}")
                doc.add_paragraph(f"Company Classification: {app.company_classification}")
                doc.add_paragraph(f"Date of Registration: {app.date_of_registration}")
                doc.add_paragraph(f"Website: {app.website or 'N/A'}")
                doc.add_paragraph(f"Turnover: {app.turnover or 'N/A'}")
                doc.add_paragraph(f"Director Count: {app.director_count}")
                
                # Project Details
                doc.add_heading('Project Details', level=2)
                doc.add_paragraph(f"Voucher Type: {app.get_voucher_type_display() if hasattr(app, 'get_voucher_type_display') else app.voucher_type}")
                doc.add_paragraph(f"Project Title: {app.project_title}")
                doc.add_paragraph(f"Product Name: {app.product_name}")
                doc.add_paragraph(f"Problem Statement: {app.problem_statement}")
                doc.add_paragraph(f"Solution Proposed: {app.solution_proposed}")
                doc.add_paragraph(f"Detailed Description: {app.detailed_description}")
                doc.add_paragraph(f"Invention/Upgrade: {app.invention_upgrade}")
                doc.add_paragraph(f"Scope: {app.scope}")
                doc.add_paragraph(f"Need: {app.need}")
                doc.add_paragraph(f"Competitive Advantage: {app.competitive_advantage}")
                doc.add_paragraph(f"Social Impact: {app.social_impact}")
                doc.add_paragraph(f"Collaboration: {app.collaboration}")
                doc.add_paragraph(f"Benefits: {app.benefits}")
                doc.add_paragraph(f"Price Advantage: {app.price_advantage}")
                doc.add_paragraph(f"Project Cost: {app.project_cost}")
                doc.add_paragraph(f"Target Dates: {app.target_dates}")
                doc.add_paragraph(f"YouTube Link: {app.youtube_link or 'N/A'}")
                
                # Project Milestone/Activity
                doc.add_heading('Project Milestone / Activity', level=2)
                doc.add_paragraph(f"Outsourcing Charges: {app.outsourcing_charges or 0} (IVP: {app.outsourcing_ivp or 0}, Applicant: {app.outsourcing_applicant or 0})")
                doc.add_paragraph(f"Raw Materials: {app.raw_materials or 0} (IVP: {app.raw_ivp or 0}, Applicant: {app.raw_applicant or 0})")
                doc.add_paragraph(f"Fabrication Charges: {app.fabrication_charges or 0} (IVP: {app.fabrication_ivp or 0}, Applicant: {app.fabrication_applicant or 0})")
                doc.add_paragraph(f"IPR Patent: {app.ipr_patent or 0} (IVP: {app.ipr_ivp or 0}, Applicant: {app.ipr_applicant or 0})")
                doc.add_paragraph(f"Testing Validation: {app.testing_validation or 0} (IVP: {app.testing_ivp or 0}, Applicant: {app.testing_applicant or 0})")
                doc.add_paragraph(f"Commercialization Support: {app.commercialization_support or 0} (IVP: {app.commercialization_ivp or 0}, Applicant: {app.commercialization_applicant or 0})")
                doc.add_paragraph(f"Other Activity 1: {app.other1_activity or 'N/A'} - Cost: {app.other1_cost or 0} (IVP: {app.other1_ivp or 0}, Applicant: {app.other1_applicant or 0})")
                doc.add_paragraph(f"Other Activity 2: {app.other2_activity or 'N/A'} - Cost: {app.other2_cost or 0} (IVP: {app.other2_ivp or 0}, Applicant: {app.other2_applicant or 0})")
                doc.add_paragraph(f"Total Estimated Cost: {app.total_estimated_cost or 0}")
                doc.add_paragraph(f"Total IVP Contribution: {app.total_ivp_contribution or 0}")
                doc.add_paragraph(f"Total Applicant Contribution: {app.total_applicant_contribution or 0}")
                
                # Knowledge Partner
                doc.add_paragraph(f"Request Knowledge Partner: {'Yes' if app.request_knowledge_partner else 'No'}")
                
                # Directors
                doc.add_heading('Director Details', level=2)
                for i, director in enumerate(directors, 1):
                    doc.add_paragraph(f"Director {i}:")
                    doc.add_paragraph(f"  Name: {director.name}")
                    doc.add_paragraph(f"  Designation: {director.designation}")
                    doc.add_paragraph(f"  DOB: {director.dob}")
                    doc.add_paragraph(f"  Mobile: {director.mobile}")
                    doc.add_paragraph(f"  Email: {director.email}")
                
                doc.add_paragraph(f"Created At: {app.created_at}")
                doc.add_page_break()
            
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = 'attachment; filename="innovation_voucher_applications.docx"'
            doc.save(response)
            return response

    return render(request, 'admin/techforms/innovationapplication/export-data.html', {})




def prior_art_search_form(request):
    try:
        logger.debug("Received request for prior_art_search_form: %s", request.method)
        states = State.objects.all()
        if request.method == 'POST':
            step = int(request.POST.get('step', 0))
            logger.debug("Processing step: %d with data: %s", step, request.POST)
            if step == 0:
                form = PriorArtSearchFormStep1(request.POST, request.FILES)
                if form.is_valid():
                    cleaned_data = form.cleaned_data
                    # Save PriorArtSearch instance
                    application_data = {k: v for k, v in cleaned_data.items() if k != 'inventors'}
                    application = PriorArtSearch(**application_data)
                    application.save()

                    # Save Inventor objects
                    content_type = ContentType.objects.get_for_model(PriorArtSearch)
                    for inventor_data in cleaned_data['inventors']:
                        Inventor.objects.create(
                            content_type=content_type,
                            object_id=application.id,
                            name=inventor_data['name'],
                            designation=inventor_data['designation'],
                            dob=inventor_data['dob'],
                            aadhaar=inventor_data['aadhaar'],
                            mobile=inventor_data['mobile'],
                            email=inventor_data['email']
                        )

                    session_data = {
                        'id': application.id,
                        'startup_name': cleaned_data['startup_name'],
                        'dpiit_number': cleaned_data['dpiit_number'],
                        'address': cleaned_data['address'],
                        'country': cleaned_data['country'],
                        'state': cleaned_data['state'].id if cleaned_data['state'] else None,
                        'district': cleaned_data['district'].id if cleaned_data['district'] else None,
                        'email': cleaned_data['email'],
                        'contact_number': cleaned_data['contact_number'],
                        'company_classification': cleaned_data['company_classification'],
                        'date_of_registration': cleaned_data['date_of_registration'].isoformat() if cleaned_data['date_of_registration'] else None,
                        'uam_certificate': cleaned_data['uam_certificate'].name if cleaned_data['uam_certificate'] else '',
                        'website': cleaned_data['website'],
                        'inventor_count': cleaned_data['inventor_count'],
                    }
                    logger.debug("Step 0 cleaned data: %s", session_data)
                    request.session['prior_art_data'] = session_data
                    request.session.modified = True
                    return JsonResponse({'status': 'success', 'next_step': 1})
                else:
                    logger.error("Prior Art Search Form Step 1 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
            elif step == 1:
                form = PriorArtSearchFormStep2(request.POST)
                if form.is_valid():
                    initial_data = request.session.get('prior_art_data', {})
                    application = PriorArtSearch.objects.get(id=initial_data['id'])
                    for key, value in form.cleaned_data.items():
                        setattr(application, key, value)
                    application.save()
                    logger.info("Prior Art Search Form submitted successfully by %s", application.startup_name)
                    del request.session['prior_art_data']
                    return JsonResponse({'status': 'success', 'message': 'Form Submitted Successfully'})
                else:
                    logger.error("Prior Art Search Form Step 2 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
        else:
            step = int(request.GET.get('step', 0))
            logger.debug("Rendering step: %d", step)
            if step == 0:
                form = PriorArtSearchFormStep1()
            elif step == 1:
                form = PriorArtSearchFormStep2()
            return render(request, 'tech_forms/prior-art-searchform.html', {'form': form, 'step': step, 'states': states})
    except Exception as e:
        logger.error("Error in prior_art_search_form: %s", str(e))
        return JsonResponse({'status': 'error', 'errors': str(e)}, status=500)



def prior_art_search_export_data(request):
    if request.method == 'POST':
        export_option = request.POST.get('export_option')
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        format_type = request.POST.get('format', 'csv')
        queryset = PriorArtSearch.objects.all()

        if export_option == 'range':
            if not (start_date and end_date):
                return render(request, 'admin/techforms/priorartsearchapplication/export-data.html', {'error': 'Please select both Start Date and End Date for date range export.'})
            queryset = queryset.filter(date_of_registration__gte=start_date, date_of_registration__lte=end_date)
        elif export_option != 'all':
            return render(request, 'admin/techforms/priorartsearchapplication/export-data.html', {'error': 'Please select a valid export option.'})

        if format_type == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="prior_art_searches.csv"'
            writer = csv.writer(response)
            writer.writerow([
                'Startup Name', 'DPIIT Number', 'Date of Registration', 'UAM Certificate', 'Website',
                'Email', 'Contact Number', 'Company Classification', 'Inventor Count',
                'Invention Title', 'Abstract', 'Scope of Invention', 'Use of Invention',
                'Functional Structure', 'Advantage', 'Key Features', 'Known Prior Art',
                'Major Assignees', 'Inventors Search', 'Remarks Info', 'Created At',
                'Address', 'Country', 'State', 'District', 'Inventors'
            ])
            for app in queryset:
                content_type = ContentType.objects.get_for_model(PriorArtSearch)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                inventors_str = '\n'.join([f"Inventor {i+1}: Name-{inv.name} - Designation-{inv.designation} - DOB-{inv.dob} - Mobile-{inv.mobile} - Email-{inv.email}" for i, inv in enumerate(inventors)])
                writer.writerow([
                    app.startup_name, app.dpiit_number, app.date_of_registration,
                    app.uam_certificate.name if app.uam_certificate else '',
                    app.website or '',
                    app.email, app.contact_number, app.company_classification, app.inventor_count,
                    app.invention_title, app.abstract, app.scope_invention, app.use_invention,
                    app.functional_structure, app.advantage, app.key_features, app.known_prior_art,
                    app.major_assignees, app.inventors_search, app.remarks_info, app.created_at,
                    app.address, app.country,
                    app.state.name if app.state else '', app.district.name if app.district else '',
                    inventors_str
                ])
            return response
        elif format_type == 'word':
            doc = Document()
            for app in queryset:
                content_type = ContentType.objects.get_for_model(PriorArtSearch)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                doc.add_paragraph(f"Startup name: {app.startup_name}")
                doc.add_paragraph(f"Dpiit number: {app.dpiit_number}")
                doc.add_paragraph(f"Address: {app.address}")
                doc.add_paragraph(f"Country: {app.country}")
                doc.add_paragraph(f"State: {app.state.name if app.state else 'N/A'}")
                doc.add_paragraph(f"District: {app.district.name if app.district else 'N/A'}")
                doc.add_paragraph(f"Email: {app.email}")
                doc.add_paragraph(f"Contact number: {app.contact_number}")
                doc.add_paragraph(f"Company classification: {app.company_classification}")
                doc.add_paragraph(f"Date of registration: {app.date_of_registration}")
                doc.add_paragraph(f"UAM certificate: {app.uam_certificate.name if app.uam_certificate else 'N/A'}")
                doc.add_paragraph(f"Website: {app.website if app.website else 'N/A'}")
                doc.add_paragraph(f"Inventor count: {app.inventor_count}")
                doc.add_paragraph(f"Invention title: {app.invention_title}")
                doc.add_paragraph(f"Abstract: {app.abstract}")
                doc.add_paragraph(f"Scope of invention: {app.scope_invention}")
                doc.add_paragraph(f"Use of invention: {app.use_invention}")
                doc.add_paragraph(f"Functional structure: {app.functional_structure}")
                doc.add_paragraph(f"Advantage: {app.advantage}")
                doc.add_paragraph(f"Key features: {app.key_features}")
                doc.add_paragraph(f"Known prior art: {app.known_prior_art if app.known_prior_art else 'N/A'}")
                doc.add_paragraph(f"Major assignees: {app.major_assignees if app.major_assignees else 'N/A'}")
                doc.add_paragraph(f"Inventors search: {app.inventors_search if app.inventors_search else 'N/A'}")
                doc.add_paragraph(f"Remarks info: {app.remarks_info if app.remarks_info else 'N/A'}")
                doc.add_paragraph(f"Created at: {app.created_at}")
                for i, inv in enumerate(inventors, 1):
                    doc.add_paragraph(f"Inventor {i} details :")
                    doc.add_paragraph(f"Name: {inv.name}")
                    doc.add_paragraph(f"Designation: {inv.designation}")
                    doc.add_paragraph(f"Dob: {inv.dob}")
                    doc.add_paragraph(f"Mobile: {inv.mobile}")
                    doc.add_paragraph(f"Email: {inv.email}")
                doc.add_paragraph('')  
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = 'attachment; filename="prior_art_searches.docx"'
            doc.save(response)
            return response

    return render(request, 'admin/techforms/priorartapplication/export-data.html', {})






import csv
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.contrib.contenttypes.models import ContentType
from docx import Document
from django.utils import timezone
from django.http import JsonResponse
from django.shortcuts import render
from .models import TrademarkApplication
from tech_repo.models import Inventor

def trademark_application_form(request):
    try:
        from logging import getLogger
        logger = getLogger(__name__)
        logger.debug("Received request for trademark_application_form: %s", request.method)
        
        if request.method == 'POST':
            step = int(request.POST.get('step', 0))
            logger.debug("Processing step: %d with data: %s", step, request.POST)
            
            if step == 0:
                form = TrademarkApplicationFormStep1(request.POST, request.FILES)
                if form.is_valid():
                    application = form.save()  # Save the initial data
                    inventor_count = form.cleaned_data.get('inventor_count', 0)
                    content_type = ContentType.objects.get_for_model(TrademarkApplication)
                    
                    # Save inventors
                    for inventor_data in form.cleaned_data.get('inventors', []):
                        Inventor.objects.create(
                            content_type=content_type,
                            object_id=application.pk,
                            name=inventor_data['name'],
                            designation=inventor_data['designation'],
                            dob=inventor_data['dob'],
                            aadhaar=inventor_data['aadhaar'],
                            mobile=inventor_data['mobile'],
                            email=inventor_data['email']
                        )
                    
                    # Store the application ID and step 0 data in the session
                    request.session['trademark_application_id'] = application.pk
                    request.session['trademark_data'] = {
                        'startup_name': application.startup_name,
                        'dpiit_number': application.dpiit_number,
                        'address': application.address,
                        'country': application.country,
                        'state': application.state.id if application.state else None,
                        'district': application.district.id if application.district else None,
                        'email': application.email,
                        'contact_number': application.contact_number,
                        'company_classification': application.company_classification,
                        'date_of_registration': application.date_of_registration.isoformat() if application.date_of_registration else None,
                        'uam_certificate': application.uam_certificate.name if application.uam_certificate else '',
                        'website': application.website or '',
                        'inventor_count': application.inventor_count,
                        'inventors': [{
                            'name': inventor_data['name'],
                            'designation': inventor_data['designation'],
                            'dob': inventor_data['dob'],
                            'aadhaar': inventor_data['aadhaar'].name if inventor_data['aadhaar'] else '',
                            'mobile': inventor_data['mobile'],
                            'email': inventor_data['email']
                        } for inventor_data in form.cleaned_data.get('inventors', [])]
                    }
                    request.session.modified = True
                    return JsonResponse({'status': 'success', 'next_step': 1})
                else:
                    logger.error("Trademark Application Form Step 1 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
            
            elif step == 1:
                # Retrieve the existing TrademarkApplication instance
                application_id = request.session.get('trademark_application_id')
                if not application_id:
                    logger.error("No application ID found in session for step 1")
                    return JsonResponse({'status': 'error', 'message': 'Session data missing. Please start over.'}, status=400)
                
                try:
                    application = TrademarkApplication.objects.get(pk=application_id)
                except TrademarkApplication.DoesNotExist:
                    logger.error("TrademarkApplication with ID %s not found", application_id)
                    return JsonResponse({'status': 'error', 'message': 'Application not found. Please start over.'}, status=400)
                
                # Update the existing instance with step 2 data
                form = TrademarkApplicationFormStep2(request.POST, request.FILES, instance=application)
                if form.is_valid():
                    application = form.save()  # Update the existing instance
                    logger.info("Trademark Application Form updated successfully for %s", application.startup_name)
                    del request.session['trademark_application_id']
                    del request.session['trademark_data']
                    return JsonResponse({'status': 'success', 'message': 'Form Submitted Successfully'})
                else:
                    logger.error("Trademark Application Form Step 2 validation failed: %s", form.errors)
                    return JsonResponse({'status': 'error', 'errors': form.errors.as_json()}, status=400)
        
        else:
            step = int(request.GET.get('step', 0))
            logger.debug("Rendering step: %d", step)
            if step == 0:
                form = TrademarkApplicationFormStep1()
            elif step == 1:
                form = TrademarkApplicationFormStep2()
            states = State.objects.all()
            return render(request, 'tech_forms/trademarkform.html', {'form': form, 'step': step, 'states': states})
    
    except Exception as e:
        logger.error("Error in trademark_application_form: %s", str(e))
        return JsonResponse({'status': 'error', 'errors': str(e)}, status=500)

def trademarkapplication_export_data(request):
    if request.method == 'POST':
        export_option = request.POST.get('export_option')
        start_date = request.POST.get('start_date')
        end_date = request.POST.get('end_date')
        format_type = request.POST.get('format', 'csv')
        queryset = TrademarkApplication.objects.all()

        if export_option == 'range':
            if not (start_date and end_date):
                return render(request, 'admin/techforms/trademarkapplication/export-data.html', {'error': 'Please select both Start Date and End Date for date range export.'})
            queryset = queryset.filter(date_of_registration__gte=start_date, date_of_registration__lte=end_date)
        elif export_option != 'all':
            return render(request, 'admin/techforms/trademarkapplication/export-data.html', {'error': 'Please select a valid export option.'})

        if format_type == 'csv':
            response = HttpResponse(content_type='text/csv')
            response['Content-Disposition'] = 'attachment; filename="trademark_applications.csv"'
            writer = csv.writer(response)
            writer.writerow(['Startup Name', 'DPIIT Number', 'Date of Registration', 'Email', 'Contact Number', 'Company Classification', 'Inventor Count', 'Slogan', 'Created At', 'Address', 'Country', 'State', 'District', 'Inventors'])
            for app in queryset:
                content_type = ContentType.objects.get_for_model(TrademarkApplication)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                inventors_str = '\n'.join([f"Inventor {i+1}: Name-{inv.name} - Designation-{inv.designation} - DOB-{inv.dob} - Mobile-{inv.mobile} - Email-{inv.email}" for i, inv in enumerate(inventors)])
                writer.writerow([app.startup_name, app.dpiit_number, app.date_of_registration, app.email, app.contact_number, app.company_classification, app.inventor_count, app.slogan, app.created_at, app.address, app.country, app.state.name if app.state else '', app.district.name if app.district else '', inventors_str])
            return response
        elif format_type == 'word':
            doc = Document()
            for app in queryset:
                content_type = ContentType.objects.get_for_model(TrademarkApplication)
                inventors = Inventor.objects.filter(content_type=content_type, object_id=app.id)
                doc.add_paragraph(f"Startup name: {app.startup_name}")
                doc.add_paragraph(f"Dpiit number: {app.dpiit_number}")
                doc.add_paragraph(f"Address: {app.address}")
                doc.add_paragraph(f"Country: {app.country}")
                doc.add_paragraph(f"State: {app.state.name if app.state else 'N/A'}")
                doc.add_paragraph(f"District: {app.district.name if app.district else 'N/A'}")
                doc.add_paragraph(f"Email: {app.email}")
                doc.add_paragraph(f"Contact number: {app.contact_number}")
                doc.add_paragraph(f"Inventor count: {app.inventor_count}")
                for i, inv in enumerate(inventors, 1):
                    doc.add_paragraph(f"Inventor {i} details :")
                    doc.add_paragraph(f"Name: {inv.name}")
                    doc.add_paragraph(f"Designation: {inv.designation}")
                    doc.add_paragraph(f"Dob: {inv.dob}")
                    doc.add_paragraph(f"Mobile: {inv.mobile}")
                    doc.add_paragraph(f"Email: {inv.email}")
                doc.add_paragraph('')  
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = 'attachment; filename="trademark_applications.docx"'
            doc.save(response)
            return response

    return render(request, 'admin/techforms/trademarkapplication/export-data.html', {})